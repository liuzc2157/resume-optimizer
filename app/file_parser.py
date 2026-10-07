"""JD/简历文件文本抽取（用于上传 JD）。

支持格式
--------
- .txt / .md          ：直接解码
- .pdf                ：pypdf 逐页抽取
- .docx               ：python-docx 抽取段落
- 其他 / 未知扩展名   ：按 UTF-8 兜底解码（失败则忽略非法字符）

注意：老版二进制 .doc（非 .docx）无法直接解析，需用户另存为 .docx 或复制文本。
"""

import io
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXT = (".txt", ".md", ".markdown", ".text", ".pdf", ".docx")


def extract_text(filename: str, data: bytes) -> str:
    """从上传文件里抽取纯文本。"""
    lower = (filename or "").lower()

    if lower.endswith((".txt", ".md", ".markdown", ".text")):
        return data.decode("utf-8", errors="ignore")

    if lower.endswith(".pdf"):
        try:
            from pypdf import PdfReader  # 延迟导入，只在传 PDF 时才需要

            reader = PdfReader(io.BytesIO(data))
            pages = [(p.extract_text() or "") for p in reader.pages]
            return "\n".join(pages).strip()
        except Exception as exc:  # noqa: BLE001
            logger.warning("[file_parser] PDF 解析失败：%s", exc)
            raise ValueError(f"PDF 解析失败：{exc}") from exc

    if lower.endswith(".docx"):
        try:
            import docx  # 延迟导入

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            # 顺带抽取表格里的文字（很多 JD 用表格排版）
            for table in d.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" ".join(cells))
            return "\n".join(parts).strip()
        except Exception as exc:  # noqa: BLE001
            logger.warning("[file_parser] DOCX 解析失败：%s", exc)
            raise ValueError(f"DOCX 解析失败：{exc}") from exc

    # 兜底：按 UTF-8 解码
    return data.decode("utf-8", errors="ignore").strip()
